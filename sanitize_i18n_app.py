"""
sanitize_i18n_app.py
i18n XLSX 清理工具 — Streamlit 图形界面

启动方式：
  python3 -m streamlit run sanitize_i18n_app.py
"""

import difflib
import io
import re as _re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import streamlit as st
import streamlit.components.v1 as _components

# 从同目录的脚本文件导入核心逻辑和常量
sys.path.insert(0, str(Path(__file__).parent))
from sanitize_i18n_xlsx import (
    TAG_INVISIBLE, TAG_NORMALIZE, TAG_NEWLINE, TAG_MIXED,
    WARN_SPACE, WARN_KEY_FORMAT, WARN_EMPTY_VALUE,
    INVISIBLE_CHARS,
    process_workbook,
)


# ── 显示用映射 ───────────────────────────────────────────────────────────────

INVISIBLE_NAMES = {
    "​": "零宽空格 U+200B",
    "﻿": "BOM U+FEFF",
    "": "行分隔符 U+2028",
    "": "段落分隔符 U+2029",
}

TAG_SHORT = {
    TAG_INVISIBLE: "不可见字符",
    TAG_NORMALIZE: "换行统一",
    TAG_NEWLINE:   "首尾换行",
    TAG_MIXED:     "首尾换行+空格",
}
WARN_SHORT = {
    WARN_SPACE:       "纯空格",
    WARN_KEY_FORMAT:  "Key格式",
    WARN_EMPTY_VALUE: "值为空",
}
TAG_ORDER  = (TAG_INVISIBLE, TAG_NORMALIZE, TAG_NEWLINE, TAG_MIXED)
WARN_ORDER = (WARN_KEY_FORMAT, WARN_EMPTY_VALUE, WARN_SPACE)


# ── 页面配置 ─────────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="i18n XLSX 清理工具",
    page_icon="🧹",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ── 全局样式 ─────────────────────────────────────────────────────────────────

st.markdown("""
<style>
/* ── 颜色变量（亮色默认） ─────────────────────────────────────── */
:root {
  --card-change-bg:   #f1f8f2;
  --card-change-bdr:  #2e7d32;
  --card-warn-bg:     #fffaf0;
  --card-warn-bdr:    #f9a825;
  --card-head-color:  #444;
  --sep-color:        #bbb;
  --col-tag-color:    #1565c0;
  --label-color:      #888;
  --pill-bg:          #eceff1;
  --pill-color:       inherit;
  --chip-change-bg:   #e8f5e9;
  --chip-change-color:#2e7d32;
  --chip-change-bdr:  #c8e6c9;
  --chip-warn-bg:     #fff3e0;
  --chip-warn-color:  #e65100;
  --chip-warn-bdr:    #ffe0b2;
  --legend-bg:        #fafafa;
  --legend-bdr:       #eee;
  --legend-color:     #666;
  --legend-b-color:   #444;
  --ratio-track-bg:   #e0e0e0;
  --radio-bdr:        #ddd;
  --radio-checked-bg: #f0f2f6;
  --radio-checked-bdr:#aaa;
  --diff-del-bg:      #ffcdd2;
  --diff-ins-bg:      #c8e6c9;
  --edited-color:     #2e7d32;
}

/* ── 深色模式覆盖 ─────────────────────────────────────────────── */
@media (prefers-color-scheme: dark) {
  :root {
    --card-change-bg:   #1a2e1c;
    --card-change-bdr:  #4caf50;
    --card-warn-bg:     #2a2210;
    --card-warn-bdr:    #ffa000;
    --card-head-color:  #ccc;
    --sep-color:        #555;
    --col-tag-color:    #64b5f6;
    --label-color:      #777;
    --pill-bg:          #2a2d2e;
    --pill-color:       #ddd;
    --chip-change-bg:   #1b3320;
    --chip-change-color:#81c784;
    --chip-change-bdr:  #2e5e34;
    --chip-warn-bg:     #2e1f00;
    --chip-warn-color:  #ffb74d;
    --chip-warn-bdr:    #5e3a00;
    --legend-bg:        #1e1e1e;
    --legend-bdr:       #333;
    --legend-color:     #999;
    --legend-b-color:   #ccc;
    --ratio-track-bg:   #333;
    --radio-bdr:        #444;
    --radio-checked-bg: #2a2d2e;
    --radio-checked-bdr:#777;
    --diff-del-bg:      #5c1e1e;
    --diff-ins-bg:      #1e3d22;
    --edited-color:     #81c784;
  }
}
[data-theme="dark"] {
  --card-change-bg:   #1a2e1c;
  --card-change-bdr:  #4caf50;
  --card-warn-bg:     #2a2210;
  --card-warn-bdr:    #ffa000;
  --card-head-color:  #ccc;
  --sep-color:        #555;
  --col-tag-color:    #64b5f6;
  --label-color:      #777;
  --pill-bg:          #2a2d2e;
  --pill-color:       #ddd;
  --chip-change-bg:   #1b3320;
  --chip-change-color:#81c784;
  --chip-change-bdr:  #2e5e34;
  --chip-warn-bg:     #2e1f00;
  --chip-warn-color:  #ffb74d;
  --chip-warn-bdr:    #5e3a00;
  --legend-bg:        #1e1e1e;
  --legend-bdr:       #333;
  --legend-color:     #999;
  --legend-b-color:   #ccc;
  --ratio-track-bg:   #333;
  --radio-bdr:        #444;
  --radio-checked-bg: #2a2d2e;
  --radio-checked-bdr:#777;
  --diff-del-bg:      #5c1e1e;
  --diff-ins-bg:      #1e3d22;
  --edited-color:     #81c784;
}

/* ── 组件样式 ───────────────────────────────────────────────────── */
.change-row { background:var(--card-change-bg); border-left:4px solid var(--card-change-bdr);
              padding:10px 14px; border-radius:6px; margin-bottom:4px; }
.warn-row   { background:var(--card-warn-bg); border-left:4px solid var(--card-warn-bdr);
              padding:10px 14px; border-radius:6px; margin-bottom:4px; }
.card-head  { font-size:0.82rem; color:var(--card-head-color); font-weight:600; margin-bottom:6px; }
.card-head .sep     { color:var(--sep-color); margin:0 4px; }
.card-head .col-tag { color:var(--col-tag-color); }
.edited-badge { color:var(--edited-color); font-size:0.75rem; margin-left:6px; }
.label      { font-size:0.74rem; color:var(--label-color); margin:4px 0 1px; }
.mono       { font-family:'SF Mono',Menlo,Consolas,monospace;
              font-size:0.9rem; word-break:break-all; line-height:1.7; }
.diff-row   { display:flex; align-items:baseline; gap:6px; margin:3px 0; }
.diff-label { font-size:0.76rem; color:var(--label-color); flex-shrink:0; white-space:nowrap; }
.str-pill   { display:inline-block; padding:3px 12px; border-radius:13px;
              background:var(--pill-bg); color:var(--pill-color);
              max-width:100%; word-break:break-all; }
.chips      { margin-top:7px; }
.chip       { display:inline-block; background:var(--chip-change-bg); color:var(--chip-change-color);
              border:1px solid var(--chip-change-bdr); border-radius:11px;
              padding:1px 10px; font-size:0.74rem; margin:2px 5px 0 0; }
.warn-row .chip { background:var(--chip-warn-bg); color:var(--chip-warn-color);
                  border-color:var(--chip-warn-bdr); }
/* 特殊字符上色 */
.real-nl    { color:#64b5f6; font-weight:bold; }
.lit-n      { color:#ff8a65; font-weight:600; }
.invis-char { color:#9e9e9e; font-weight:bold; }
.space-char { color:#757575; }
/* 图例 */
.legend     { background:var(--legend-bg); border:1px solid var(--legend-bdr); border-radius:6px;
              padding:7px 14px; font-size:0.8rem; color:var(--legend-color); margin-bottom:14px; }
.legend b   { color:var(--legend-b-color); }
.legend .real-nl,.legend .lit-n,.legend .invis-char,.legend .space-char { font-size:1rem; margin:0 2px; }
/* 总结比例条 */
.ratio-row  { margin-bottom:9px; }
.ratio-head { display:flex; justify-content:space-between;
              font-size:0.85rem; margin-bottom:2px; color:var(--card-head-color); }
.ratio-track{ background:var(--ratio-track-bg); border-radius:4px; height:14px; overflow:hidden; }
.ratio-fill { height:100%; border-radius:4px; }
/* radio 横排 Tab */
div[data-testid="stRadio"] > div { display:flex; gap:8px; flex-wrap:wrap; }
div[data-testid="stRadio"] label {
    border:1px solid var(--radio-bdr); border-radius:6px; padding:4px 14px;
    cursor:pointer; font-size:0.9rem;
}
div[data-testid="stRadio"] label:has(input:checked) {
    background:var(--radio-checked-bg); border-color:var(--radio-checked-bdr); font-weight:600;
}
</style>
""", unsafe_allow_html=True)


# ── HTML 显示格式化 ───────────────────────────────────────────────────────────

def _esc(s) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _fmt_html(val: str, show_spaces: bool = False) -> str:
    segments = val.split("\n")
    processed = []
    for seg in segments:
        s = seg.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if show_spaces:
            s = s.replace(" ", "\x00")
        for ch in INVISIBLE_CHARS:
            name = INVISIBLE_NAMES.get(ch, f"U+{ord(ch):04X}")
            s = s.replace(ch, f'<span class="invis-char" title="{name}">⬚</span>')
        s = s.replace("\\n", '<span class="lit-n">\\n</span>')
        if show_spaces:
            s = s.replace("\x00", '<span class="space-char">␣</span>')
        processed.append(s)
    return '<span class="real-nl">↵</span>'.join(processed)


def _diff_html(old: str, new: str) -> tuple[str, str]:
    matcher = difflib.SequenceMatcher(None, old, new, autojunk=False)
    old_parts: list[str] = []
    new_parts: list[str] = []
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        old_seg = _fmt_html(old[i1:i2])
        new_seg = _fmt_html(new[j1:j2])
        if op == "equal":
            old_parts.append(old_seg)
            new_parts.append(new_seg)
        elif op == "replace":
            old_parts.append(f'<mark style="background:var(--diff-del-bg);border-radius:2px">{old_seg}</mark>')
            new_parts.append(f'<mark style="background:var(--diff-ins-bg);border-radius:2px">{new_seg}</mark>')
        elif op == "delete":
            old_parts.append(f'<mark style="background:var(--diff-del-bg);border-radius:2px">{old_seg}</mark>')
        elif op == "insert":
            new_parts.append(f'<mark style="background:var(--diff-ins-bg);border-radius:2px">{new_seg}</mark>')
    return "".join(old_parts), "".join(new_parts)


def render_legend() -> None:
    st.markdown(
        '<div class="legend"><b>图例</b>　'
        '<span class="real-nl">↵</span>真实换行　'
        '<span class="lit-n">\\n</span>字面量　'
        '<span class="space-char">␣</span>空格　'
        '<span class="invis-char">⬚</span>不可见字符（悬停看详情）　'
        '<mark style="background:var(--diff-ins-bg);border-radius:2px">新增</mark>　'
        '<mark style="background:var(--diff-del-bg);border-radius:2px">删除</mark>'
        '</div>',
        unsafe_allow_html=True,
    )


# ── 文案辅助 ─────────────────────────────────────────────────────────────────

def _change_reasons(old: str, tags: set) -> list:
    out = []
    if TAG_INVISIBLE in tags:
        names = [INVISIBLE_NAMES.get(ch, f"U+{ord(ch):04X}")
                 for ch in INVISIBLE_CHARS if ch in old]
        out.append("删除 " + "、".join(names) if names else "删除不可见字符")
    if TAG_NORMALIZE in tags:
        out.append("字面量 \\n → 真实换行")
    if TAG_NEWLINE in tags:
        out.append("去除首尾换行")
    return out


def _humanize_warn(val: str, detail: str) -> str:
    if not (detail.startswith("整个值") or
            detail.startswith("行首") or
            detail.startswith("行尾")):
        return detail
    stripped_l = val.lstrip("\n \t")
    stripped_r = val.rstrip("\n \t")
    leading  = val[: len(val) - len(stripped_l)]
    trailing = val[len(stripped_r):]
    if detail.startswith("整个值") or not stripped_l:
        seg, loc = val, "整个单元格"
    elif detail.startswith("行首"):
        seg, loc = leading, "行首"
    else:
        seg, loc = trailing, "行尾"
    n_nl = seg.count("\n")
    n_sp = len(seg) - n_nl
    bits = []
    if n_nl: bits.append(f"换行符 ×{n_nl}")
    if n_sp: bits.append(f"空格 ×{n_sp}")
    desc = " + ".join(bits) if bits else "空白"
    if loc == "整个单元格":
        return f"整个单元格仅由空白组成（{desc}）"
    return f"{loc}有 {desc}"


# ── 筛选条 ───────────────────────────────────────────────────────────────────

def _filter_bar(prefix: str, type_opts: list,
                sheets: list, cols: list) -> tuple:
    c1, c2 = st.columns([3, 2])
    with c1:
        sel_type = st.segmented_control(
            "类型", ["全部"] + type_opts, default="全部",
            key=f"flt_{prefix}_type", label_visibility="collapsed",
        ) or "全部"
    with c2:
        query = st.text_input(
            "搜索", placeholder="🔍 搜索 key 或内容",
            key=f"flt_{prefix}_q", label_visibility="collapsed",
        )
    sel_sheet = sel_col = "全部"
    if len(sheets) > 1 or len(cols) > 1:
        c3, c4 = st.columns(2)
        if len(sheets) > 1:
            sel_sheet = c3.selectbox("工作表", ["全部"] + sheets, key=f"flt_{prefix}_sheet")
        if len(cols) > 1:
            sel_col = c4.selectbox("语言列", ["全部"] + cols, key=f"flt_{prefix}_col")
    return sel_type, query, sel_sheet, sel_col


def _inject_textarea_autoresize() -> None:
    """注入 JS：让页面内所有 textarea 随内容自动撑高，不出现滚动条。"""
    _components.html("""
    <script>
    (function(){
        var doc = window.parent.document;
        function fit(ta){
            ta.style.overflowY = 'hidden';
            ta.style.height    = 'auto';
            ta.style.height    = ta.scrollHeight + 'px';
            var lc = doc.getElementById('lc-current');
            if (lc) {
                var n = (ta.value.match(/\n/g) || []).length + 1;
                lc.textContent = n + ' 行';
            }
        }
        function attachAll(){
            doc.querySelectorAll('textarea').forEach(function(ta){
                if (ta.dataset.autofit) return;
                ta.dataset.autofit = '1';
                ta.addEventListener('input', function(){ fit(ta); });
                fit(ta);
            });
        }
        attachAll();
        if (!window.parent.__autofitObserver){
            window.parent.__autofitObserver =
                new MutationObserver(attachAll);
            window.parent.__autofitObserver.observe(doc.body,
                {childList: true, subtree: true});
        }
    })();
    </script>
    """, height=0, scrolling=False)


def _reset_filters() -> None:
    for k in list(st.session_state.keys()):
        if k.startswith("flt_"):
            del st.session_state[k]
    st.session_state.cell_edits   = {}
    st.session_state.editing_cell = None


# ── 编辑状态管理 ─────────────────────────────────────────────────────────────

def _uid(sheet: str, row: int, col_idx: int) -> str:
    """生成按钮/输入框的唯一 key（不含特殊字符）。"""
    safe = _re.sub(r"[^a-zA-Z0-9]", "", sheet)[:12]
    return f"{safe}_{row}_{col_idx}"


def _apply_edit(sheet: str, row: int, col_idx: int, new_value: str) -> None:
    """将用户输入写入活跃 workbook 并重新序列化下载内容。"""
    wb  = st.session_state.results["live_wb"]
    ck  = (sheet, row, col_idx)
    # 首次编辑时保存当前值作为撤销基准
    base = st.session_state.cell_edits.get(ck, {}).get(
        "base", wb[sheet].cell(row=row, column=col_idx).value or "")
    wb[sheet].cell(row=row, column=col_idx).value = new_value
    st.session_state.cell_edits[ck] = {"current": new_value, "base": base}
    st.session_state.editing_cell   = None
    buf = io.BytesIO()
    wb.save(buf)
    st.session_state.results["output"] = buf.getvalue()


def _undo_edit(sheet: str, row: int, col_idx: int) -> None:
    """撤销用户编辑，恢复至编辑前的自动处理结果。"""
    ck    = (sheet, row, col_idx)
    entry = st.session_state.cell_edits.pop(ck, None)
    if entry is None:
        return
    wb = st.session_state.results["live_wb"]
    wb[sheet].cell(row=row, column=col_idx).value = entry["base"]
    buf = io.BytesIO()
    wb.save(buf)
    st.session_state.results["output"] = buf.getvalue()


def _render_edit_ui(sheet: str, row: int, col_idx: int,
                    edited_label: str = "修正字符：",
                    btn_label:    str = "✏️ 手动修改") -> None:
    """
    在卡片下方渲染编辑控件，三种状态：
      · 正常：显示编辑按钮
      · 编辑中：显示 textarea + 确认/取消
      · 已修改：显示已修改值 + 撤销/重新编辑
    """
    ck        = (sheet, row, col_idx)
    uid       = _uid(sheet, row, col_idx)
    is_editing = st.session_state.get("editing_cell") == ck
    edit_info  = st.session_state.get("cell_edits", {}).get(ck)

    if is_editing:
        wb      = st.session_state.results["live_wb"]
        cur_val = wb[sheet].cell(row=row, column=col_idx).value or ""
        ta_key  = f"ta_{uid}"
        st.text_area(
            "编辑内容", value=cur_val, key=ta_key,
            label_visibility="collapsed", height=40,
        )
        init_lines = cur_val.count("\n") + 1
        st.markdown(
            f'<div id="lc-current" style="font-size:0.72rem;color:var(--label-color);'
            f'text-align:left;margin:-6px 0 4px">{init_lines} 行</div>',
            unsafe_allow_html=True,
        )
        c1, c2, _ = st.columns([1, 1, 5])
        if c1.button("✅ 确认", key=f"ok_{uid}"):
            _apply_edit(sheet, row, col_idx, st.session_state[ta_key])
            st.rerun()
        if c2.button("✖ 取消", key=f"cancel_{uid}"):
            st.session_state.editing_cell = None
            st.rerun()

    elif edit_info:
        # 已修改态：展示修正后的值（使用符号系统显示）
        st.markdown(
            f'<div class="diff-row" style="margin-top:2px">'
            f'<span class="diff-label">{_esc(edited_label)}</span>'
            f'<span class="str-pill mono">{_fmt_html(edit_info["current"])}</span>'
            f'</div>',
            unsafe_allow_html=True,
        )
        c1, c2, _ = st.columns([1, 1, 5])
        if c1.button("↩️ 撤销", key=f"undo_{uid}"):
            _undo_edit(sheet, row, col_idx)
            st.rerun()
        if c2.button("✏️ 改", key=f"reedit_{uid}"):
            st.session_state.editing_cell = ck
            st.rerun()

    else:
        if st.button(btn_label, key=f"edit_{uid}"):
            st.session_state.editing_cell = ck
            st.rerun()


# ── 渲染：变更 ───────────────────────────────────────────────────────────────

def render_changes(changes: list) -> None:
    render_legend()

    present   = set().union(*(tags for *_, tags in changes)) if changes else set()
    type_opts = [TAG_SHORT[t] for t in TAG_ORDER if t in present]
    sheets    = list(dict.fromkeys(s for s, *_ in changes))
    cols      = list(dict.fromkeys(c for _, _, _, _, c, *_ in changes))

    sel_type, query, sel_sheet, sel_col = _filter_bar("c", type_opts, sheets, cols)
    q = query.lower().strip()

    filtered = []
    for item in changes:
        sheet, row, col_idx, key, col, old, new, tags = item
        if sel_type != "全部" and sel_type not in {TAG_SHORT.get(t, t) for t in tags}:
            continue
        if sel_sheet != "全部" and sheet != sel_sheet:
            continue
        if sel_col != "全部" and col != sel_col:
            continue
        if q and q not in f"{key}\n{old}\n{new}".lower():
            continue
        filtered.append(item)

    st.caption(f"显示 {len(filtered)} / {len(changes)} 处变更"
               + ("（无匹配项）" if not filtered else ""))

    groups: dict = {}
    for item in filtered:
        groups.setdefault(item[3], []).append(item)   # key at index 3

    cell_edits = st.session_state.get("cell_edits", {})

    for key, items in groups.items():
        label = key if key else "（无 key）"
        with st.expander(f"🔑 {label}　{len(items)} 处", expanded=True):
            for sheet, row, col_idx, _key, col, old, new, tags in items:
                ck           = (sheet, row, col_idx)
                edited_badge = '<span class="edited-badge">✅ 已修改</span>' if ck in cell_edits else ''
                old_d, new_d = _diff_html(old, new)
                chips        = "".join(f'<span class="chip">{_esc(r)}</span>'
                                       for r in _change_reasons(old, tags))
                st.markdown(
                    f'<div class="change-row">'
                    f'<div class="card-head">📄 {_esc(sheet)}<span class="sep">·</span>'
                    f'第 {row} 行 <span class="col-tag">[{_esc(col)}]</span>'
                    f'{edited_badge}</div>'
                    f'<div class="diff-row"><span class="diff-label">原始字符：</span>'
                    f'<span class="str-pill mono">{old_d}</span></div>'
                    f'<div class="diff-row"><span class="diff-label">修正字符：</span>'
                    f'<span class="str-pill mono">{new_d}</span></div>'
                    f'<div class="chips">{chips}</div>'
                    f'</div>',
                    unsafe_allow_html=True,
                )
                _render_edit_ui(sheet, row, col_idx,
                                edited_label="调整后：",
                                btn_label="✏️ 调整")


# ── 渲染：警告 ───────────────────────────────────────────────────────────────

def render_warnings(warn_items: list) -> None:
    render_legend()

    present = set()
    for *_, msgs in warn_items:
        present |= {wt for wt, _ in msgs}
    type_opts = [WARN_SHORT[t] for t in WARN_ORDER if t in present]
    sheets    = list(dict.fromkeys(s for s, *_ in warn_items))
    cols      = list(dict.fromkeys(c for _, _, _, _, c, *_ in warn_items))

    sel_type, query, sel_sheet, sel_col = _filter_bar("w", type_opts, sheets, cols)
    q = query.lower().strip()

    filtered = []
    for item in warn_items:
        sheet, row, col_idx, key, col, val, msgs = item
        types_here = {WARN_SHORT.get(wt, wt) for wt, _ in msgs}
        if sel_type != "全部" and sel_type not in types_here:
            continue
        if sel_sheet != "全部" and sheet != sel_sheet:
            continue
        if sel_col != "全部" and col != sel_col:
            continue
        if q and q not in f"{key}\n{val}".lower():
            continue
        filtered.append(item)

    st.caption(f"显示 {len(filtered)} / {len(warn_items)} 处警告"
               + ("（无匹配项）" if not filtered else ""))

    groups: dict = {}
    for item in filtered:
        groups.setdefault(item[3], []).append(item)   # key at index 3

    cell_edits = st.session_state.get("cell_edits", {})

    for key, items in groups.items():
        label = key if key else "（无 key）"
        with st.expander(f"🔑 {label}　{len(items)} 处", expanded=True):
            for sheet, row, col_idx, _key, col, val, msgs in items:
                ck           = (sheet, row, col_idx)
                edited_badge = '<span class="edited-badge">✅ 已修改</span>' if ck in cell_edits else ''
                chips        = "".join(
                    f'<span class="chip">⚠ {_esc(_humanize_warn(val, d))}</span>'
                    for _, d in msgs)
                val_section  = (
                    f'<div class="label">当前值</div>'
                    f'<div class="mono">{_fmt_html(val, show_spaces=True)}</div>'
                ) if val else ""
                st.markdown(
                    f'<div class="warn-row">'
                    f'<div class="card-head">📄 {_esc(sheet)}<span class="sep">·</span>'
                    f'第 {row} 行 <span class="col-tag">[{_esc(col)}]</span>'
                    f'{edited_badge}</div>'
                    f'{val_section}'
                    f'<div class="chips">{chips}</div>'
                    f'</div>',
                    unsafe_allow_html=True,
                )
                _render_edit_ui(sheet, row, col_idx,
                                edited_label="修正字符：",
                                btn_label="✏️ 手动修改")


# ── 渲染：总结 ───────────────────────────────────────────────────────────────

def _ratio_bar(label: str, count: int, total: int, color: str) -> None:
    pct = (count / total * 100) if total else 0
    st.markdown(
        f'<div class="ratio-row">'
        f'<div class="ratio-head"><span>{_esc(label)}</span><span><b>{count}</b></span></div>'
        f'<div class="ratio-track"><div class="ratio-fill" '
        f'style="width:{pct:.1f}%;background:{color}"></div></div></div>',
        unsafe_allow_html=True,
    )


def render_summary(changes: list, warn_items: list) -> None:
    tag_counts:       dict = defaultdict(int)
    warn_type_counts: dict = defaultdict(int)
    sheet_c:          dict = defaultdict(int)
    sheet_w:          dict = defaultdict(int)

    for *_, tags in changes:
        for tag in tags:
            tag_counts[tag] += 1
    for sheet, *_ in changes:
        sheet_c[sheet] += 1
    for sheet, _, _, _, _, _, msgs in warn_items:
        sheet_w[sheet] += 1
        for wt, _ in msgs:
            warn_type_counts[wt] += 1

    n_edited    = len(st.session_state.get("cell_edits", {}))
    cleanup     = st.session_state.results.get("cleanup", {})
    del_rows    = sum(v["rows"] for v in cleanup.values())
    del_cols    = sum(v["cols"] for v in cleanup.values())

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("✅ 自动修改", len(changes))
    c2.metric("⚠️ 警告", len(warn_items))
    c3.metric("✏️ 手动修改", n_edited)
    c4.metric("🗑️ 删除空行/列", f"{del_rows} 行 / {del_cols} 列")
    st.divider()

    col_a, col_b = st.columns(2)
    with col_a:
        st.markdown("**变更类型**")
        tc_total = max(sum(tag_counts.values()), 1)
        if any(tag_counts.get(t) for t in TAG_ORDER):
            for tag in TAG_ORDER:
                n = tag_counts.get(tag, 0)
                if n:
                    _ratio_bar(tag, n, tc_total, "#66bb6a")
        else:
            st.caption("无变更")
    with col_b:
        st.markdown("**警告类型**")
        wc_total = max(sum(warn_type_counts.values()), 1)
        if any(warn_type_counts.get(t) for t in WARN_ORDER):
            for wt in WARN_ORDER:
                n = warn_type_counts.get(wt, 0)
                if n:
                    _ratio_bar(wt, n, wc_total, "#ffa726")
        else:
            st.caption("无警告")

    st.divider()
    st.markdown("**按工作表（自动修改 / 警告）**")
    all_sheets = list(dict.fromkeys(
        [s for s, *_ in changes] + [s for s, *_ in warn_items]
    ))
    if all_sheets:
        st.markdown(
            '<table style="width:100%;border-collapse:collapse;font-size:0.9rem">'
            '<thead><tr style="background:var(--radio-checked-bg)">'
            '<th style="text-align:left;padding:6px 10px">工作表</th>'
            '<th style="text-align:center;padding:6px 10px">自动修改</th>'
            '<th style="text-align:center;padding:6px 10px">警告</th>'
            '</tr></thead><tbody>'
            + "".join(
                f'<tr style="border-top:1px solid #e0e0e0">'
                f'<td style="padding:5px 10px">{_esc(s)}</td>'
                f'<td style="text-align:center;padding:5px 10px;color:#2e7d32">'
                f'<b>{sheet_c.get(s, 0)}</b></td>'
                f'<td style="text-align:center;padding:5px 10px;color:#e65100">'
                f'<b>{sheet_w.get(s, 0)}</b></td></tr>'
                for s in all_sheets
            )
            + "</tbody></table>",
            unsafe_allow_html=True,
        )
    else:
        st.caption("没有任何修改或警告。")


# ── 主界面 ───────────────────────────────────────────────────────────────────

def main() -> None:

    with st.sidebar:
        st.header("📋 处理规则")
        st.markdown("""
**步骤 1　统一为真实换行**
- 删除不可见字符
  （零宽空格、BOM、行/段分隔符）
- 行尾 `\\r\\n` / `\\r` → 换行
- 字面量 `\\n` → 真实换行

**步骤 2　处理首尾**
- 仅换行 → 自动去除
- 换行＋空格混合 → 自动去除
- 仅空格 → ⚠ 警告

**校验（只警告，不修改）**
- Key 格式：只允许小写字母、数字和下划线
- 语言值：每个语言列不能为空

> 处理后可在页面内手动修改任意单元格，
> 下载即最终稿，无需二次对比。
""")

    st.title("🧹 i18n XLSX 清理工具")
    st.caption("上传翻译文件，自动清理格式问题，在页面内修正警告，下载即最终稿")

    uploaded = st.file_uploader(
        "将 xlsx 文件拖放到此处，或点击选择",
        type=["xlsx"],
        label_visibility="collapsed",
    )

    if uploaded is None:
        st.session_state.pop("results", None)
        st.info("📂 请上传一个 xlsx 文件开始处理")
        return

    st.success(f"已选择：**{uploaded.name}**（{len(uploaded.getvalue()) / 1024:.1f} KB）")

    if st.button("🚀 开始处理", type="primary", use_container_width=True):
        with st.spinner("处理中，请稍候…"):
            wb = openpyxl.load_workbook(io.BytesIO(uploaded.getvalue()))
            changes, warn_items, cleanup = process_workbook(wb)
            output = io.BytesIO()
            wb.save(output)

        st.session_state.results = {
            "changes":    changes,
            "warn_items": warn_items,
            "cleanup":    cleanup,
            "output":     output.getvalue(),
            "filename":   uploaded.name,
            "live_wb":    wb,
        }
        _reset_filters()

    if "results" not in st.session_state:
        return

    _inject_textarea_autoresize()   # textarea 动态高度（每次 rerun 确保注入）

    r          = st.session_state.results
    changes    = r["changes"]
    warn_items = r["warn_items"]
    n_edited   = len(st.session_state.get("cell_edits", {}))

    # 下载按钮：始终反映最新状态（含手动编辑）
    dl_label = (f"⬇️ 下载处理结果（含 {n_edited} 处手动修改）"
                if n_edited else "⬇️ 下载处理结果")
    st.download_button(
        label=dl_label,
        data=r["output"],
        file_name=Path(r["filename"]).stem + "_sanitized.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    st.divider()

    TAB_SUMMARY  = "📊 总结"
    TAB_CHANGES  = f"✅ 变更  {len(changes)} 处"
    TAB_WARNINGS = f"⚠️ 警告  {len(warn_items)} 处"

    active_tab = st.radio(
        "tab_nav",
        options=[TAB_SUMMARY, TAB_CHANGES, TAB_WARNINGS],
        horizontal=True,
        label_visibility="collapsed",
        key="active_tab",
    )

    st.divider()

    if active_tab == TAB_SUMMARY:
        render_summary(changes, warn_items)
    elif active_tab == TAB_CHANGES:
        if changes:
            render_changes(changes)
        else:
            st.success("无需修改，文件已是规范格式。")
    else:
        if warn_items:
            st.info("以下条目需人工确认，可直接在此修改后下载。")
            render_warnings(warn_items)
        else:
            st.success("没有需要人工确认的警告。")


if __name__ == "__main__":
    main()
